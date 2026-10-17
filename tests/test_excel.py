import unittest,json,zipfile
from pathlib import Path
from datetime import datetime,time
import openpyxl
P=Path(__file__).resolve().parents[1]

class ExcelCompatible(unittest.TestCase):
    def test_todas_las_filas_tipos_y_contrato(self):
        d=json.loads((P/'datos_compatibles/fuentes_sinteticas.json').read_text())
        for c,p in d['proyectos'].items():
            path=P/f'datos_compatibles/GHTOP_{c}_sintetico.xlsx'
            wb=openpyxl.load_workbook(path,read_only=True,data_only=False);ws=wb[p['hoja']];rows=list(ws.iter_rows(values_only=True));idx=0 if c=='A' else 1
            self.assertEqual(list(rows[idx][:len(p['cabeceras'])]),p['cabeceras']);self.assertEqual(len(rows[idx]),39 if c=='A' else 21)
            self.assertEqual(len(rows)-idx-1,len(p['registros']))
            for got,exp in zip(rows[idx+1:],p['registros']):
                for key,v in zip(p['cabeceras'],got):
                    e=exp[key]
                    if key in ['Fecha','FECHA']:self.assertIsInstance(v,datetime);self.assertEqual(v.date().isoformat(),e)
                    elif key in ['Hora inicio','Hora termino','HORA INICIO','HORA FIN'] and e is not None:self.assertIsInstance(v,time);self.assertEqual(v.isoformat(),e)
                    elif isinstance(e,(int,float)):self.assertIsInstance(v,(int,float));self.assertAlmostEqual(v,e)
                    else:self.assertEqual(v,e)
            self.assertEqual(wb.sheetnames,[p['hoja'],'CalendarioOperacion','Canchas','LEEME'])
            cal=list(wb['CalendarioOperacion'].iter_rows(values_only=True));self.assertEqual(len(cal)-1,len(p['calendario']))
            self.assertTrue(all(isinstance(r[0],datetime) and isinstance(r[1],bool) for r in cal[1:]));wb.close()
            with zipfile.ZipFile(path) as z:
                self.assertFalse(any('externalLink' in n for n in z.namelist()))
                self.assertFalse(any(b'<f>' in z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/') and n.endswith('.xml')))

if __name__=='__main__':unittest.main()
